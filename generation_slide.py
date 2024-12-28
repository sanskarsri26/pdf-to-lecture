import google.generativeai as genai
from pptx import Presentation
from pptx.util import Pt

# Configure the API key for Google Generative AI
genai.configure(api_key="AIzaSyCSqsv69biM6pCAkPGEDh9XRM5WpVBraf4")

# Define the prompt
prompt = """
Create a professional lecture on Natural Language Processing (NLP), designed for a 1-hour session. The lecture should include two key elements for each slide:

1. **Speaker Notes**: Detailed, concise, and engaging script that the speaker will use to explain each topic while the corresponding slide is displayed. The script should provide enough depth to keep the audience interested but should be simple to follow.

2. **Slide Content**: Text-based content for the slides, structured in bullet points or short statements. The slide content should be visually appealing, avoid overwhelming detail, and support the speaker's script.

The presentation should cover the following topics:
- **Introduction to NLP**: Definition, importance, and real-world applications.
- **How NLP Works**: Breaking down key components like tokenization, parsing, and semantics.
- **Core Techniques**: Overview of machine learning, deep learning, and transformer models in NLP.
- **Real-World Challenges**: Handling ambiguity, context, and large datasets.
- **Future of NLP**: Trends like multilingual models, real-time processing, and ethical considerations.

For each slide, provide:
- Slide Title.
- Slide Content (bullet points).
- Speaker Notes (detailed script).

The session should last approximately 1 hour. Allocate 2–5 minutes per slide based on content complexity. Where appropriate, suggest visuals, animations, or diagrams to enhance the slides.
"""

# Generate content using the Generative AI API
model = genai.GenerativeModel("gemini-1.5-flash")
response = model.generate_content(prompt)

# Extract the response text and print for debugging
response_text = (
    response.candidates[0].content.parts[0].text
)  # Ensure correct attribute access
print("Gemini API Response Text:")
print(response_text)

# Parse the response into structured slide data
response_data = {"slides": []}
slides = response_text.split("\n\n**Slide")  # Split based on slide markers
for slide in slides:
    lines = slide.strip().split("\n")
    slide_info = {"title": "", "content": [], "notes": ""}
    current_section = None
    for line in lines:
        line = line.strip()
        if line.startswith("Title:"):
            slide_info["title"] = line.replace("Title:", "").strip()
        elif line.startswith("* **Slide Content:**"):
            current_section = "content"
        elif line.startswith("* **Speaker Notes:**"):
            current_section = "notes"
        elif line.startswith("    *") or line.startswith("    "):
            if current_section == "content":
                slide_info["content"].append(line.strip("* ").strip())
            elif current_section == "notes":
                slide_info["notes"] += line.strip("* ").strip() + " "
    if slide_info["title"]:  # Only add slides with a valid title
        response_data["slides"].append(slide_info)

# Debug parsed slide data
print("\nParsed Slide Data:")
for slide in response_data["slides"]:
    print(slide)

# Initialize PowerPoint presentation
prs = Presentation()


# Function to add a slide with content and speaker notes
def add_slide_with_notes(title, content, notes):
    slide_layout = prs.slide_layouts[1]  # Title and content layout
    slide = prs.slides.add_slide(slide_layout)
    slide.shapes.title.text = title

    # Add bullet points for slide content
    text_frame = slide.placeholders[1].text_frame
    text_frame.clear()
    for point in content:
        p = text_frame.add_paragraph()
        p.text = point
        p.font.size = Pt(18)

    # Add speaker notes
    notes_slide = slide.notes_slide
    notes_text_frame = notes_slide.notes_text_frame
    notes_text_frame.text = notes


# Add slides dynamically based on parsed response
for slide_data in response_data["slides"]:
    add_slide_with_notes(
        slide_data["title"], slide_data["content"], slide_data["notes"]
    )

# Save the PowerPoint presentation
filename = "NLP_Lecture_With_Notes.pptx"
prs.save(filename)
print(f"Presentation created and saved as '{filename}'.")
